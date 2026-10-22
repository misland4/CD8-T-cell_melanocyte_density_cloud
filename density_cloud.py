"""Plot CD8+ T-cell / melanocyte counts from a supplied or legacy Excel sheet.

Author information will be updated upon publication of the paper.
"""
import argparse
from pathlib import Path
from zipfile import BadZipFile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from scipy.stats import gaussian_kde


ROOT = Path(__file__).resolve().parent

DARK_GREEN = (53 / 255, 143 / 255, 83 / 255)
LIGHT_GREEN = (149 / 255, 197 / 255, 111 / 255)


def read_points(ws, start_row, end_row, x_col, y_col):
    points = []
    columns = f"{get_column_letter(x_col)} and {get_column_letter(y_col)}"
    for row in range(start_row, end_row + 1):
        x = ws.cell(row, x_col).value
        y = ws.cell(row, y_col).value
        if x is None and y is None:
            continue
        if x is None or y is None:
            raise ValueError(f"Both counts are required at row {row} (columns {columns}).")
        try:
            pair = np.array([float(x), float(y)])
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Non-numeric count at row {row} (columns {columns}).") from exc
        if not np.isfinite(pair).all() or (pair < 0).any():
            raise ValueError(f"Counts must be finite and non-negative at row {row}.")
        points.append(pair)
    result = np.asarray(points, dtype=float).reshape(-1, 2)
    if len(result) < 3 or np.linalg.matrix_rank(result - result.mean(axis=0)) < 2:
        raise ValueError(
            f"Rows {start_row}-{end_row} need at least three non-collinear points for 2D KDE."
        )
    return result


def read_groups(ws):
    """Use one supplied B/C group, or the original two D/E groups."""
    x_header = str(ws.cell(1, 2).value or "").strip().casefold()
    y_header = str(ws.cell(1, 3).value or "").strip().casefold()
    if x_header == "cd8+ t cells" and y_header in {"melanocyte", "melanocytes"}:
        return read_points(ws, 2, ws.max_row, 2, 3), None
    return read_points(ws, 17, 23, 4, 5), read_points(ws, 24, 30, 4, 5)


def kde_density(points, xi, yi, bandwidth):
    grid = np.vstack([xi.ravel(), yi.ravel()])
    try:
        kde = gaussian_kde(points.T, bw_method=bandwidth)
    except np.linalg.LinAlgError as exc:
        raise ValueError("KDE covariance is singular; supply non-collinear counts.") from exc
    z = kde(grid).reshape(xi.shape)
    peak = z.max()
    if not np.isfinite(peak) or peak <= 0:
        raise ValueError("No finite KDE density inside the plot limits; check the count ranges.")
    return z / peak


def cloud_cmap(name, rgb, max_alpha):
    return LinearSegmentedColormap.from_list(
        name,
        [
            (0.00, (rgb[0], rgb[1], rgb[2], 0.00)),
            (0.25, (rgb[0], rgb[1], rgb[2], max_alpha * 0.18)),
            (0.70, (rgb[0], rgb[1], rgb[2], max_alpha * 0.58)),
            (1.00, (rgb[0], rgb[1], rgb[2], max_alpha)),
        ],
    )


def draw_vector_cloud(ax, points, rgb, xi, yi, bandwidth, cutoff, max_alpha):
    z = kde_density(points, xi, yi, bandwidth)
    z = np.ma.masked_where(z < cutoff, z)
    levels = np.linspace(cutoff, 1.0, 18)
    ax.contourf(xi, yi, z, levels=levels, cmap=cloud_cmap("cloud", rgb, max_alpha), antialiased=True, zorder=1)
    ax.scatter(points[:, 0], points[:, 1], s=34, color=rgb, edgecolors="white", linewidths=0.45, zorder=3)


def create_figure(dark_points, light_points=None):
    """Keep the original group colors, KDE bandwidths, and fixed axis ranges."""
    plt.rcParams["pdf.fonttype"] = 42
    plt.rcParams["svg.fonttype"] = "path"
    xi, yi = np.mgrid[0:601:4, 0:501:4]

    fig = plt.figure(figsize=(6.4, 4.8))
    ax = fig.add_axes([0.18, 0.18, 0.77, 0.76])
    draw_vector_cloud(ax, dark_points, DARK_GREEN, xi, yi, bandwidth=0.58, cutoff=0.04, max_alpha=0.54)
    if light_points is not None:
        draw_vector_cloud(ax, light_points, LIGHT_GREEN, xi, yi, bandwidth=0.64, cutoff=0.04, max_alpha=0.56)

    ax.set_xlim(0, 600)
    ax.set_ylim(0, 500)
    ax.set_xlabel("No. of CD8+ T cells per scan view", fontsize=10.5, fontweight="bold", labelpad=7)
    ax.set_ylabel("No. of melanocytes per scan view", fontsize=10.5, fontweight="bold", labelpad=5)
    ax.set_xticks(np.arange(0, 601, 200))
    ax.set_yticks(np.arange(0, 501, 100))
    ax.set_xticks(np.arange(0, 601, 20), minor=True)
    ax.tick_params(axis="both", which="major", direction="out", length=6, width=1.4, labelsize=8.5, pad=1.5)
    ax.tick_params(axis="x", which="minor", direction="out", length=3.5, width=1.1)
    for label in ax.get_xticklabels() + ax.get_yticklabels():
        label.set_fontweight("bold")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["bottom"].set_linewidth(1.4)
    ax.spines["left"].set_linewidth(1.4)

    return fig, ax


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True,
                        help="Excel .xlsx: B/C with row-1 headers, or legacy D/E rows 17-30.")
    parser.add_argument("--sheet", help="Worksheet name (default: active worksheet).")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "output")
    args = parser.parse_args(argv)
    try:
        wb = load_workbook(args.input, data_only=True)
        try:
            if args.sheet and args.sheet not in wb.sheetnames:
                raise ValueError(f"Worksheet not found: {args.sheet}")
            ws = wb[args.sheet] if args.sheet else wb.active
            dark_points, light_points = read_groups(ws)
        finally:
            wb.close()
        points = dark_points if light_points is None else np.vstack([dark_points, light_points])
        if (points[:, 0] > 600).any() or (points[:, 1] > 500).any():
            raise ValueError("Counts exceed the original axis limits (x: 0-600, y: 0-500).")
        fig, _ = create_figure(dark_points, light_points)
        try:
            args.output_dir.mkdir(parents=True, exist_ok=True)
            for extension in ("pdf", "png", "svg"):
                output = args.output_dir / f"density_cloud.{extension}"
                fig.savefig(output, dpi=300, bbox_inches="tight", facecolor="white")
                print(output)
        finally:
            plt.close(fig)
    except (OSError, ValueError, BadZipFile, InvalidFileException) as exc:
        parser.error(str(exc))


if __name__ == "__main__":
    main()
