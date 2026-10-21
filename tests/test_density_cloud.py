"""Exercise portable CLI output and reject unusable scientific input."""
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook


SCRIPT = Path(__file__).resolve().parents[1] / "density_cloud.py"


class DensityCloudTests(unittest.TestCase):
    def run_cli(self, folder, points):
        workbook = Workbook()
        sheet = workbook.active
        for row, pair in enumerate(points, start=17):
            sheet.cell(row, 4, pair[0])
            sheet.cell(row, 5, pair[1])
        source = folder / "input.xlsx"
        workbook.save(source)
        workbook.close()
        return subprocess.run(
            [sys.executable, str(SCRIPT), "--input", str(source),
             "--output-dir", str(folder / "figures")],
            cwd=folder, env={**os.environ, "MPLBACKEND": "Agg", "PYTHONIOENCODING": "utf-8"},
            capture_output=True, text=True, encoding="utf-8",
        )

    def test_workbook_from_another_directory_produces_three_formats(self):
        points = [(20, 40), (60, 20), (80, 75), (110, 55),
                  (150, 95), (180, 60), (210, 130),
                  (250, 230), (290, 200), (330, 270), (370, 220),
                  (410, 320), (450, 280), (490, 350)]
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            result = self.run_cli(folder, points)
            self.assertEqual(result.returncode, 0, result.stderr)
            for extension in ("png", "pdf", "svg"):
                output = folder / "figures" / f"density_cloud.{extension}"
                self.assertTrue(output.is_file(), str(output))
                self.assertGreater(output.stat().st_size, 100)

    def test_collinear_points_report_an_input_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = self.run_cli(Path(tmp), [(n, n) for n in range(1, 15)])
            self.assertEqual(result.returncode, 2, result.stderr)
            self.assertIn("non-collinear", result.stderr)
            self.assertNotIn("Traceback", result.stderr)

    def test_half_empty_pair_is_not_silently_dropped(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = self.run_cli(Path(tmp), [(20, None)] * 14)
            self.assertEqual(result.returncode, 2, result.stderr)
            self.assertIn("row 17", result.stderr)
            self.assertNotIn("Traceback", result.stderr)

    def test_non_excel_input_reports_an_input_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "wrong.csv"
            source.write_text("x,y\n1,2\n", encoding="utf-8")
            result = subprocess.run(
                [sys.executable, str(SCRIPT), "--input", str(source)],
                env={**os.environ, "PYTHONIOENCODING": "utf-8"},
                capture_output=True, text=True, encoding="utf-8",
            )
            self.assertEqual(result.returncode, 2, result.stderr)
            self.assertNotIn("Traceback", result.stderr)

    def test_single_group_supplied_workbook_produces_three_formats(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            result = subprocess.run(
                [sys.executable, str(SCRIPT), "--input",
                 str(SCRIPT.parent / "examples" / "research_density.xlsx"),
                 "--output-dir", str(folder)],
                cwd=folder, env={**os.environ, "PYTHONIOENCODING": "utf-8"},
                capture_output=True, text=True, encoding="utf-8",
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            for extension in ("png", "pdf", "svg"):
                self.assertGreater((folder / f"density_cloud.{extension}").stat().st_size, 100)


if __name__ == "__main__":
    unittest.main()
